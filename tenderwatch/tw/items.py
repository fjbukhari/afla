"""Read the list of items a tender asks for from its bidding document (PDF or Excel),
so tenders can be matched line by line against the JB Scientific catalogue."""
import io
import re

_CLAUSE = re.compile(r"\b(shall|bidders?|tenderers?|must|will be|procuring|agency|evaluation|security|clause|"
                     r"signature|stamp|affidavit|warranty|delivery period|payment|penalt|blacklist|rules?|"
                     r"submitted|submission|opening|envelope|earnest|validity|conditions?|eligib)\w*", re.I)
_NUM = re.compile(r"^\s*(\d{1,3})\s*[.):\-]?\s+(.+?)\s*$")
_DESC_HEAD = re.compile(r"descr|specification|particular|item name|name of item|items?\b|product|material|test name", re.I)
_BAD_HEAD = re.compile(r"^(s\.?\s*no|sr|#|no\.?|qty|quantity|unit|uom|rate|price|amount|total|remarks?)\b", re.I)


def _is_item(text):
    t = text.strip()
    if len(t) < 6 or len(t) > 260 or not re.search(r"[A-Za-z]{3}", t):
        return False
    if _CLAUSE.search(t):
        return False
    return True


def items_from_text(text, limit=300):
    """Numbered lines that form a run 1, 2, 3... are taken as the item list."""
    lines = [l for l in (text or "").splitlines() if l.strip()]
    runs, cur, expect = [], [], None
    for line in lines:
        m = _NUM.match(line)
        if not m:
            continue
        n, body = int(m.group(1)), m.group(2)
        if expect is not None and n == expect:
            cur.append(body)
            expect = n + 1
        elif n == 1:
            if len(cur) >= 3:
                runs.append(cur)
            cur, expect = [body], 2
        else:
            continue
    if len(cur) >= 3:
        runs.append(cur)
    out = []
    for run in runs:
        good = [b for b in run if _is_item(b)]
        if len(good) >= max(3, len(run) // 2):  # mostly product-like lines, not numbered clauses
            out.extend(good)
    seen, uniq = set(), []
    for b in out:
        b = re.sub(r"\s{2,}", " ", b)
        if b.lower() not in seen:
            seen.add(b.lower())
            uniq.append(b)
    return uniq[:limit]


def items_from_pdf(data, limit=300):
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    text = "\n".join((p.extract_text() or "") for p in reader.pages[:60])
    return items_from_text(text, limit)


def items_from_rows(rows, limit=300):
    """Spreadsheet rows (lists of cells): find the header row and its description column."""
    for hi, row in enumerate(rows[:15]):
        cells = [str(c or "").strip() for c in row]
        col = next((i for i, c in enumerate(cells) if _DESC_HEAD.search(c) and not _BAD_HEAD.match(c)), None)
        if col is None:
            continue
        vals = [str(r[col]).strip() for r in rows[hi + 1:] if col < len(r) and r[col] is not None]
        vals = [v for v in vals if _is_item(v) and not re.fullmatch(r"[\d.,\s]+", v)]
        if len(vals) >= 1:
            return list(dict.fromkeys(vals))[:limit]
    return []


def items_from_xlsx(data, limit=300):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(max_row=2000, values_only=True)]
        found = items_from_rows(rows, limit)
        if found:
            return found
    return []


def items_from_document(data, url="", content_type=""):
    kind = (content_type or "").lower() + " " + url.lower()
    if data[:4] == b"%PDF" or "pdf" in kind:
        return items_from_pdf(data)
    if data[:2] == b"PK" or "spreadsheet" in kind or re.search(r"\.xlsx(\?|$)", url, re.I):
        return items_from_xlsx(data)
    return []


def fetch_items(browser, con, log=print, max_docs=40, max_bytes=15_000_000):
    """Download documents of open relevant tenders not yet read, and store their item lists."""
    con.execute("CREATE TABLE IF NOT EXISTS tender_items (id TEXT, n INTEGER, text TEXT, PRIMARY KEY (id, n))")
    con.execute("CREATE TABLE IF NOT EXISTS item_docs (id TEXT PRIMARY KEY, checked TEXT, status TEXT)")
    todo = con.execute(
        "SELECT t.id, t.doc_url FROM tenders t LEFT JOIN item_docs d ON d.id = t.id "
        "WHERE t.relevant = 1 AND t.dup_of IS NULL AND t.doc_url != '' AND d.id IS NULL "
        "AND (t.closing IS NULL OR t.closing >= date('now')) ORDER BY t.closing LIMIT ?", (max_docs,)).fetchall()
    from .db import now
    got = 0
    for row in todo:
        status = "no items found"
        try:
            r = browser.ctx.request.get(row["doc_url"], timeout=60000)
            body = r.body() if r.ok else b""
            if not r.ok:
                status = f"HTTP {r.status}"
            elif len(body) > max_bytes:
                status = "document too large"
            else:
                items = items_from_document(body, row["doc_url"], r.headers.get("content-type", ""))
                if items:
                    con.execute("DELETE FROM tender_items WHERE id = ?", (row["id"],))
                    con.executemany("INSERT INTO tender_items VALUES (?,?,?)",
                                    [(row["id"], i + 1, t) for i, t in enumerate(items)])
                    status = f"{len(items)} items"
                    got += 1
        except Exception as e:  # a bad document must not stop the run
            status = f"error: {type(e).__name__}"
        con.execute("INSERT OR REPLACE INTO item_docs VALUES (?,?,?)", (row["id"], now(), status))
        con.commit()
    if todo:
        log(f"Tender documents: read {len(todo)}, item lists found in {got}")
    return got


def items_for(con):
    con.execute("CREATE TABLE IF NOT EXISTS tender_items (id TEXT, n INTEGER, text TEXT, PRIMARY KEY (id, n))")
    out = {}
    for r in con.execute("SELECT id, text FROM tender_items ORDER BY id, n"):
        out.setdefault(r["id"], []).append(r["text"])
    return out
